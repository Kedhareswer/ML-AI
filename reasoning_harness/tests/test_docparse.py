"""docparse: exact structure parsers, layout analysis, and (when weights exist) the OCR model."""
import datetime as dt
import random
import tempfile
import unittest
from pathlib import Path

import numpy as np

from docparse.extract import DEFAULT_MODEL, extract, to_markdown, to_text
from docparse.ocr.charset import CHARS, decode_greedy, encode, normalize
from docparse.ocr.metrics import cer, summary
from docparse.ocr.page import analyze, group
from docparse.synth_pages import make_page


class StubReader:
    """Stands in for the network: returns the box geometry as text, so layout can be tested alone."""

    def read(self, images, batch=32):
        return [(f"w{a.shape[1]}", 1.0) for a in images]


class CharsetTests(unittest.TestCase):
    def test_normalize_folds_typography(self):
        self.assertEqual(normalize("“Café” – naïve…"), '"Cafe" - naive...')

    def test_encode_decode_roundtrip_with_ctc_blanks(self):
        h, e, l, o = encode("helo")
        self.assertEqual(decode_greedy([0, h, h, 0, e, l, 0, l, o, o]), "hello")
        self.assertEqual(decode_greedy([l, l]), "l")  # repeats need a blank between them
        self.assertEqual(len(CHARS), len(set(CHARS)))

    def test_metrics(self):
        self.assertEqual(cer("abc", "abc"), 0)
        self.assertAlmostEqual(cer("abd", "abc"), 1 / 3)
        s = summary(["a b", "c"], ["a b", "d"])
        self.assertEqual(s["line_exact"], 0.5)


class StructureTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_docx_headings_lists_tables_merged_cells(self):
        import docx
        d = docx.Document()
        d.sections[0].header.paragraphs[0].text = "Header text"
        d.add_heading("Main title", 1)
        d.add_paragraph("Body paragraph.")
        d.add_paragraph("Bullet", style="List Bullet")
        d.add_paragraph("Nested", style="List Bullet 2")
        t = d.add_table(rows=2, cols=3)
        t.cell(0, 0).merge(t.cell(0, 1)).text = "Wide"
        t.cell(0, 2).text = "C"
        for c, v in enumerate("xyz"):
            t.cell(1, c).text = v
        p = d.add_paragraph()
        p.add_run("Total").bold = True
        d.save(self.dir / "a.docx")
        doc = extract(self.dir / "a.docx")
        types = [(b["type"], b.get("text")) for b in doc["blocks"] if b["type"] != "table"]
        self.assertIn(("header", "Header text"), types)
        self.assertIn(("heading", "Main title"), types)
        self.assertIn(("paragraph", "Body paragraph."), types)
        items = [b for b in doc["blocks"] if b["type"] == "list_item"]
        self.assertEqual([(b["text"], b.get("level", 0)) for b in items], [("Bullet", 0), ("Nested", 1)])
        table = next(b for b in doc["blocks"] if b["type"] == "table")
        self.assertEqual(table["rows"], [["Wide", "", "C"], ["x", "y", "z"]])
        self.assertIn("| Wide |  | C |", to_markdown(doc))

    def test_xlsx_values_formulas_regions_dates(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["Date", "Qty", "Price", "Total"])
        ws.append([dt.date(2024, 2, 1), 2, 1.5, "=B2*C2"])
        ws["G1"], ws["G2"] = "Note", "apart"
        ws.merge_cells("A5:B5")
        ws["A5"] = "Footer"
        wb.save(self.dir / "b.xlsx")
        doc = extract(self.dir / "b.xlsx")
        tables = [b for b in doc["blocks"] if b["type"] == "table"]
        self.assertEqual(tables[0]["range"], "A1:D2")
        self.assertEqual(tables[0]["rows"][1], ["2024-02-01", 2, 1.5, "=B2*C2"])  # no cached value -> formula
        self.assertEqual([t["range"] for t in tables[1:]], ["G1:G2", "A5:A5"])
        sheet = doc["sheets"][0]
        self.assertEqual(sheet["formulas"], {"D2": "=B2*C2"})
        self.assertEqual(sheet["merged"], ["A5:B5"])

    def test_csv_sniffs_delimiter(self):
        (self.dir / "c.csv").write_text("a;b\n1;2,5\n", encoding="utf-8")
        doc = extract(self.dir / "c.csv")
        self.assertEqual(doc["blocks"][0]["rows"], [["a", "b"], ["1", "2,5"]])

    def test_pdf_text_layer_headings_paragraphs_tables(self):
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError:
            self.skipTest("reportlab not installed")
        st = getSampleStyleSheet()
        body = ("The quarterly figures were reviewed by the committee and found to be in line with "
                "the forecast published at the start of the year, with small differences in two regions.")
        t = Table([["Region", "Units", "Amount"], ["North", "120", "1,250.00"], ["South", "98", "980.50"]])
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.black)]))
        SimpleDocTemplate(str(self.dir / "d.pdf"), pagesize=A4).build(
            [Paragraph("Annual Review", st["Title"]), Paragraph(body, st["BodyText"]), Spacer(1, 12),
             Paragraph(body.replace("quarterly", "monthly"), st["BodyText"]), Spacer(1, 12), t])
        doc = extract(self.dir / "d.pdf", ocr="never")
        kinds = [b["type"] for b in doc["blocks"]]
        self.assertEqual(kinds, ["heading", "paragraph", "paragraph", "table"])
        self.assertEqual(doc["blocks"][0]["text"], "Annual Review")
        self.assertEqual(doc["blocks"][1]["text"], body)
        self.assertEqual(doc["blocks"][3]["rows"][1], ["North", "120", "1,250.00"])
        self.assertTrue(all(b["source"] == "text-layer" for b in doc["blocks"]))


class PdfDetailTests(unittest.TestCase):
    def test_inline_bold_italic_and_superscript(self):
        from docparse.extract import _styled_text, plain

        def w(text, x0, x1, font="Times-Roman", size=10, bottom=100):
            return {"text": text, "x0": x0, "x1": x1, "top": bottom - size, "bottom": bottom, "size": size,
                    "fontname": font}
        ws = [w("Water", 0, 30), w("is", 33, 40, "Times-Bold"), w("H", 43, 50), w("2", 50.2, 54, size=6, bottom=102),
              w("O", 54.3, 61), w("ref", 64, 80, "Times-Italic"), w("1", 80.2, 83, size=6, bottom=96)]
        out = _styled_text(ws, 10)
        self.assertEqual(out, "Water **is** H<sub>2</sub>O *ref*<sup>1</sup>")
        self.assertEqual(plain(out), "Water is H2O ref1")

    def test_page_furniture(self):
        from docparse.extract import mark_page_furniture
        blocks = [{"type": "paragraph", "text": "Annual Report 2024", "bbox": [50, 20, 300, 35]},
                  {"type": "heading", "text": "Title", "bbox": [50, 40, 300, 60], "level": 1},
                  {"type": "paragraph", "text": "Body", "bbox": [50, 300, 300, 320]},
                  {"type": "paragraph", "text": "Page 3 of 10", "bbox": [250, 770, 330, 785]}]
        mark_page_furniture(blocks, 792)
        self.assertEqual([b["type"] for b in blocks], ["header", "heading", "paragraph", "footer"])
        self.assertTrue(blocks[3].get("page_number"))

    def test_stacked_table_rows_are_split(self):
        from docparse.extract import _split_stacked_rows
        rows = [["Class", "Rate"], ["8833\n8837", "1.08\n0.87"], ["Total", "2"]]
        self.assertEqual(_split_stacked_rows(rows), [["Class", "Rate"], ["8833", "1.08"], ["8837", "0.87"],
                                                     ["Total", "2"]])


class LayoutTests(unittest.TestCase):
    def test_synthetic_page_layout_with_stub_reader(self):
        img, truth = make_page(3, "test")
        blocks, meta = analyze(img, StubReader())
        kinds = [b["type"] for b in blocks]
        self.assertEqual(kinds.count("heading"), sum(b["type"] == "heading" for b in truth))
        tables = [b for b in blocks if b["type"] == "table"]
        want = sorted((len(b["rows"]), len(b["rows"][0])) for b in truth if b["type"] == "table")
        self.assertEqual(sorted((len(t["rows"]), len(t["rows"][0])) for t in tables), want)
        # Two-column section: the left column is read before the right column.
        paras = [b for b in blocks if b["type"] == "paragraph"]
        cols = [p for p in paras if p["bbox"][1] > 290 and p["bbox"][1] < 320]
        self.assertEqual(len(cols), 2)
        self.assertLess(cols[0]["bbox"][0], cols[1]["bbox"][0])

    def test_scanned_page_is_deskewed(self):
        img, _ = make_page(11, "test", scan=True)
        img = img.rotate(2.5, fillcolor=255)
        _, meta = analyze(img, StubReader())
        self.assertGreater(abs(meta["deskew_degrees"]), 1.5)

    def test_group_merges_paragraph_lines_and_keeps_code_breaks(self):
        def line(y, x1, text):
            return {"kind": "line", "box": (y, y + 20, 100, x1), "segs": [(y, y + 20, 100, x1)],
                    "parts": [(text, 0.9)], "xh": 12}
        prose = [line(0, 900, "one two three four"), line(28, 905, "five six seven eight"), line(56, 300, "nine")]
        out = group(prose, 20)
        self.assertEqual([b["text"] for b in out], ["one two three four five six seven eight nine"])
        code = [line(0, 300, "import os"), line(28, 340, "x = os.sep"), line(56, 380, "print(x)"),
                line(84, 900, "# a long comment line that runs the full width")]
        self.assertEqual(group(code, 20)[0]["text"].split("\n")[:3], ["import os", "x = os.sep", "print(x)"])


@unittest.skipUnless(DEFAULT_MODEL.exists(), "trained OCR weights not present")
class ModelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from docparse.extract import recognizer
        cls.rec = recognizer()

    def test_heldout_font_lines(self):
        from docparse.ocr.render import dataset
        ims, txt = dataset(200, 424242, "test", degrade=False)
        res = summary([p for p, _ in self.rec.read(ims)], txt)
        self.assertLess(res["cer"], 0.05, res)

    def test_image_file_end_to_end(self):
        img, truth = make_page(21, "test")
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "page.png"
            img.save(p)
            doc = extract(p)
        text = to_text(doc)
        ref = " ".join(b["text"] for b in truth if b["type"] == "heading")
        words = [w for w in ref.split() if len(w) > 3]
        found = sum(w in text for w in words)
        self.assertGreaterEqual(found, 0.8 * len(words), (text[:500], words))
        self.assertTrue(all(b["source"] == "ocr" for b in doc["blocks"]))

    def test_scanned_pdf_page_uses_ocr(self):
        img, _ = make_page(22, "test", scan=True)
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "scan.pdf"
            img.convert("RGB").save(p, "PDF", resolution=150)
            doc = extract(p)
        self.assertTrue(doc["blocks"])
        self.assertTrue(all(b["source"] == "ocr" for b in doc["blocks"]))
        x0, y0, x1, y1 = doc["blocks"][0]["bbox"]
        self.assertLess(x1, 600)  # bbox converted to PDF points (A4 width is 595)


if __name__ == "__main__":
    random.seed(0)
    np.random.seed(0)
    unittest.main()
